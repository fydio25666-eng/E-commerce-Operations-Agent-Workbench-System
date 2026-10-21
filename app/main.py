from io import BytesIO
import asyncio
import logging
import os
import re
import time
from typing import Any
from pathlib import Path
from zipfile import BadZipFile

import xlrd
import xlwt
from openpyxl import Workbook, load_workbook
from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, ValidationError
from openai import OpenAIError
from langgraph.types import Command
import uuid
from contextlib import asynccontextmanager
from sqlalchemy import delete, select

from app.database import SessionLocal, create_tables
from app.models import ProductRecord, ProductVersion
from app.schemas.csv_row import ProductInput
from app.schemas.copy import GeneratedCopy
from app.graph.state import ProductState
from app.graph.workflow import workflow
from app.services.checker import check_violations
from app.services.llm_service import generate_copy_with_metadata
from app.utils.logger import configure_logging

configure_logging()
logger = logging.getLogger(__name__)

@asynccontextmanager
async def lifespan(_: FastAPI):
    await create_tables()
    yield


app = FastAPI(title="电商运营工作台", lifespan=lifespan)
frontend_origins = [
    origin.strip()
    for origin in os.getenv("FRONTEND_ORIGINS", "").split(",")
    if origin.strip()
]
if frontend_origins:
    app.add_middleware(
        CORSMiddleware,
        allow_origins=frontend_origins,
        allow_credentials=False,
        allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
        allow_headers=["*"],
    )
products: list[ProductInput] = []
task_progress: dict[str, dict[str, Any]] = {}
task_controls: dict[str, asyncio.Event] = {}
product_threads: dict[str, str] = {}


def _session_factory():
    if SessionLocal is None:
        raise HTTPException(status_code=503, detail="未配置 DATABASE_URL，请检查 .env")
    return SessionLocal


async def _save_product(
    session,
    product_id: str,
    *,
    status: str | None = None,
    retry_count: int | None = None,
    error_log: str | None = None,
    process_duration: float | None = None,
) -> ProductRecord:
    record = await session.get(ProductRecord, product_id)
    if record is None:
        record = ProductRecord(product_id=product_id, status=status or "待审核")
        session.add(record)
    if status is not None:
        record.status = status
    if retry_count is not None:
        record.retry_count = retry_count
    record.error_log = error_log
    if process_duration is not None:
        record.process_duration = process_duration
    return record


def _add_version(session, product_id: str, version_type: str, content: dict[str, Any]) -> None:
    session.add(
        ProductVersion(product_id=product_id, version_type=version_type, copy_json=content)
    )


async def _record_failure(
    product_id: str, error: str, process_duration: float | None = None
) -> None:
    if SessionLocal is None:
        return
    async with SessionLocal() as session:
        await _save_product(
            session,
            product_id,
            status="失败",
            error_log=error,
            process_duration=process_duration,
        )
        await session.commit()


async def _process_upload_task(task_id: str, imported: list[ProductInput]) -> None:
    progress = task_progress[task_id]
    semaphore = asyncio.Semaphore(3)
    resume_event = task_controls[task_id]

    async def process_one(product: ProductInput) -> None:
        thread_id = str(uuid.uuid4())
        config = {"configurable": {"thread_id": thread_id}}
        initial_state: ProductState = {
            "item_code": product.product_id,
            "raw_data": product.model_dump(by_alias=True),
            "generated_copy": None,
            "violations": [],
            "retry_count": 0,
            "review_status": "pending",
            "model_source": "unknown",
            "model_retry_count": 0,
            "generation_status": "processing",
            "generation_error": None,
        }
        async with semaphore:
            await resume_event.wait()
            started_at = time.perf_counter()
            model_source = "unknown"
            model_retry_count = 0
            violation_count = 0
            process_duration = 0.0
            product_status = "FAILED"
            try:
                await workflow.ainvoke(initial_state, config=config)
                snapshot = await workflow.aget_state(config)
                state = snapshot.values
                model_source = state.get("model_source", "unknown")
                model_retry_count = state.get("model_retry_count", 0)
                violation_count = len(state.get("violations", []))
                process_duration = time.perf_counter() - started_at
                is_fallback = state.get("generation_status") == "FALLBACK_DRAFT"
                product_status = "FALLBACK_DRAFT" if is_fallback else "PENDING_REVIEW"
                generated = state.get("generated_copy")
                if hasattr(generated, "model_dump"):
                    generated = generated.model_dump()
                async with _session_factory()() as session:
                    await _save_product(
                        session,
                        product.product_id,
                        status="FALLBACK_DRAFT" if is_fallback else "待审核",
                        retry_count=state.get("retry_count", 0) + model_retry_count,
                        error_log=state.get("generation_error"),
                        process_duration=process_duration,
                    )
                    _add_version(
                        session,
                        product.product_id,
                        "FALLBACK_DRAFT" if is_fallback else "AI_GENERATED",
                        generated,
                    )
                    await session.commit()
                progress["success_count"] += 1
                completed_product = {
                    "item_code": product.product_id,
                    "raw": product.model_dump(by_alias=True),
                    "copy": generated,
                    "violations": state.get("violations", []),
                    "thread_id": thread_id,
                    "status": product_status,
                    "model_source": model_source,
                    "retry_count": model_retry_count,
                    "process_duration": round(process_duration, 3),
                    "error_log": state.get("generation_error"),
                }
                for index, queued_product in enumerate(progress["products"]):
                    if queued_product["item_code"] == product.product_id:
                        progress["products"][index] = completed_product
                        break
                progress["updates"].append(completed_product)
            except Exception as exc:
                process_duration = time.perf_counter() - started_at
                logger.exception("批量商品处理失败 item_code=%s", product.product_id)
                try:
                    await _record_failure(product.product_id, str(exc), process_duration)
                except Exception:
                    pass
                progress["errors"].append({"item_code": product.product_id, "error": str(exc)})
                for queued_product in progress["products"]:
                    if queued_product["item_code"] == product.product_id:
                        queued_product["status"] = "FAILED"
                        queued_product["error"] = str(exc)
                        break
                progress["updates"].append({
                    "item_code": product.product_id,
                    "status": "FAILED",
                    "error": str(exc),
                })
            finally:
                process_duration = time.perf_counter() - started_at
                logger.info(
                    "product_generation item_code=%s model_source=%s duration_seconds=%.3f "
                    "retry_count=%s violation_count=%s status=%s",
                    product.product_id,
                    model_source,
                    process_duration,
                    model_retry_count,
                    violation_count,
                    product_status,
                )
                progress["processed_products"] += 1
                progress["processed_count"] = progress["processed_products"]

    await asyncio.gather(*(process_one(product) for product in imported))
    progress["status"] = "completed"
    resume_event.set()


@app.get("/api/tasks/{task_id}/progress")
async def get_task_progress(task_id: str, cursor: int = 0) -> dict[str, Any]:
    progress = task_progress.get(task_id)
    if progress is None:
        raise HTTPException(status_code=404, detail="任务不存在或服务已重启")
    return {
        "task_id": task_id,
        "status": progress["status"],
        "total_products": progress["total_products"],
        "processed_products": progress["processed_products"],
        "processed_count": progress["processed_count"],
        "success_count": progress["success_count"],
        "updates": progress["updates"][max(cursor, 0):],
        "next_cursor": len(progress["updates"]),
        "errors": progress["errors"],
        "validation_errors": progress["validation_errors"],
    }


@app.post("/api/tasks/{task_id}/pause")
async def pause_task(task_id: str) -> dict[str, str]:
    progress = task_progress.get(task_id)
    if progress is None:
        raise HTTPException(status_code=404, detail="任务不存在或服务已重启")
    if progress["status"] == "completed":
        raise HTTPException(status_code=409, detail="任务已完成")
    task_controls[task_id].clear()
    progress["status"] = "paused"
    return {"task_id": task_id, "status": "paused"}


@app.post("/api/tasks/{task_id}/resume")
async def resume_task(task_id: str) -> dict[str, str]:
    progress = task_progress.get(task_id)
    if progress is None:
        raise HTTPException(status_code=404, detail="任务不存在或服务已重启")
    if progress["status"] == "completed":
        raise HTTPException(status_code=409, detail="任务已完成")
    task_controls[task_id].set()
    progress["status"] = "processing"
    return {"task_id": task_id, "status": "processing"}


class ProductReview(BaseModel):
    action: str
    copy: GeneratedCopy | None = None


class ProductBatchDelete(BaseModel):
    item_codes: list[str]


def _latest_product_data(product: ProductRecord, versions: list[ProductVersion]) -> dict[str, Any]:
    ordered = sorted(versions, key=lambda version: version.id)
    raw = next((version.copy_json for version in ordered if version.version_type == "RAW"), {})
    copy = next(
        (version.copy_json for version in reversed(ordered) if version.version_type != "RAW"),
        {},
    )
    status_map = {
        "已通过": "APPROVED",
        "待审核": "PENDING_REVIEW",
        "RAW": "RAW",
        "失败": "FAILED",
        "AI_GENERATED": "AI_GENERATED",
        "FALLBACK_DRAFT": "FALLBACK_DRAFT",
    }
    return {
        "item_code": product.product_id,
        "raw": raw,
        "copy": copy,
        "product_name": raw.get("产品名称", ""),
        "material": raw.get("产品特性和材质", ""),
        "dimensions": raw.get("产品的尺寸", ""),
        "price": raw.get("产品价格", ""),
        "weight": raw.get("产品的重量", ""),
        "product_features": raw.get("产品特性", ""),
        "title": copy.get("title", ""),
        "bullet_points": copy.get("bullet_points", []),
        "keywords": copy.get("keywords", []),
        "unverified_fields": copy.get("unverified_fields", []),
        "status": status_map.get(product.status, product.status),
        "retry_count": product.retry_count,
        "error_log": product.error_log,
        "process_duration": product.process_duration,
    }


async def _product_summary(session) -> list[dict[str, Any]]:
    records = (await session.scalars(select(ProductRecord).order_by(ProductRecord.product_id))).all()
    if not records:
        return []
    ids = [record.product_id for record in records]
    versions = (
        await session.scalars(
            select(ProductVersion)
            .where(ProductVersion.product_id.in_(ids))
            .order_by(ProductVersion.product_id, ProductVersion.id)
        )
    ).all()
    by_product: dict[str, list[ProductVersion]] = {}
    for version in versions:
        by_product.setdefault(version.product_id, []).append(version)
    result = [_latest_product_data(record, by_product.get(record.product_id, [])) for record in records]
    active_threads = {
        item["item_code"]: item.get("thread_id")
        for task in task_progress.values()
        for item in task.get("products", [])
        if item.get("thread_id")
    }
    active_threads.update(product_threads)
    for item in result:
        item["thread_id"] = active_threads.get(item["item_code"])
    return result


@app.get("/api/products/summary")
async def products_summary() -> dict[str, Any]:
    async with _session_factory()() as session:
        return {"products": await _product_summary(session)}


@app.delete("/api/products/{item_code}")
async def delete_product(item_code: str) -> dict[str, Any]:
    async with _session_factory()() as session:
        record = await session.get(ProductRecord, item_code)
        if record is None:
            raise HTTPException(status_code=404, detail="商品不存在")
        await session.execute(delete(ProductVersion).where(ProductVersion.product_id == item_code))
        await session.delete(record)
        await session.commit()
    return {"deleted_count": 1, "item_codes": [item_code]}


@app.post("/api/products/batch-delete")
async def delete_products_batch(request: ProductBatchDelete) -> dict[str, Any]:
    item_codes = list(dict.fromkeys(code.strip() for code in request.item_codes if code.strip()))
    if not item_codes:
        raise HTTPException(status_code=422, detail="item_codes 不能为空")
    async with _session_factory()() as session:
        existing = (
            await session.scalars(select(ProductRecord.product_id).where(ProductRecord.product_id.in_(item_codes)))
        ).all()
        if existing:
            await session.execute(delete(ProductVersion).where(ProductVersion.product_id.in_(existing)))
            await session.execute(delete(ProductRecord).where(ProductRecord.product_id.in_(existing)))
            await session.commit()
    return {"deleted_count": len(existing), "item_codes": list(existing)}


@app.get("/api/products/export")
async def export_products_xlsx() -> StreamingResponse:
    async with _session_factory()() as session:
        rows = await _product_summary(session)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品全景"
    headers = [
        "商品编号", "原始名称", "原始规格参数", "AI生成标题", "卖点1", "卖点2",
        "卖点3", "卖点4", "卖点5", "关键词", "审核状态",
    ]
    sheet.append(headers)
    for row in rows:
        raw = row["raw"]
        specs = "；".join(
            f"{label}：{raw.get(key, '')}"
            for label, key in (
                ("特性", "产品特性"), ("材质", "产品特性和材质"), ("价格", "产品价格"),
                ("重量", "产品的重量"), ("尺寸", "产品的尺寸"),
            )
            if raw.get(key)
        )
        bullets = (row["bullet_points"] + [""] * 5)[:5]
        sheet.append([
            row["item_code"], row["product_name"], specs, row["title"], *bullets,
            "、".join(row["keywords"]), row["status"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="products.xlsx"'},
    )


@app.post("/api/products/process")
async def process_product(product: ProductInput) -> dict[str, Any]:
    thread_id = str(uuid.uuid4())
    product_threads[product.product_id] = thread_id
    config = {"configurable": {"thread_id": thread_id}}
    initial_state: ProductState = {
        "item_code": product.product_id,
        "raw_data": product.model_dump(by_alias=True),
        "generated_copy": None,
        "violations": [],
        "retry_count": 0,
        "review_status": "pending",
        "model_source": "unknown",
        "model_retry_count": 0,
        "generation_status": "processing",
        "generation_error": None,
    }
    started_at = time.perf_counter()
    try:
        async with _session_factory()() as session:
            await _save_product(session, product.product_id, status="待审核", retry_count=0)
            _add_version(session, product.product_id, "RAW", product.model_dump(by_alias=True))
            await session.commit()
        await workflow.ainvoke(initial_state, config=config)
        snapshot = await workflow.aget_state(config)
    except HTTPException:
        raise
    except RuntimeError as exc:
        await _record_failure(product.product_id, str(exc), time.perf_counter() - started_at)
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except ValidationError as exc:
        await _record_failure(product.product_id, str(exc), time.perf_counter() - started_at)
        raise HTTPException(status_code=502, detail={"message": "模型返回结构无效", "errors": exc.errors()}) from exc
    except OpenAIError as exc:
        await _record_failure(product.product_id, str(exc), time.perf_counter() - started_at)
        raise HTTPException(status_code=502, detail=f"大模型请求失败：{exc}") from exc

    state = snapshot.values
    process_duration = time.perf_counter() - started_at
    model_source = state.get("model_source", "unknown")
    model_retry_count = state.get("model_retry_count", 0)
    violation_count = len(state.get("violations", []))
    is_fallback = state.get("generation_status") == "FALLBACK_DRAFT"
    generated = state.get("generated_copy")
    if hasattr(generated, "model_dump"):
        generated = generated.model_dump()
    async with _session_factory()() as session:
        await _save_product(
            session,
            product.product_id,
            status="FALLBACK_DRAFT" if is_fallback else "待审核",
            retry_count=state.get("retry_count", 0) + model_retry_count,
            error_log=state.get("generation_error"),
            process_duration=process_duration,
        )
        _add_version(
            session,
            product.product_id,
            "FALLBACK_DRAFT" if is_fallback else "AI_GENERATED",
            generated,
        )
        await session.commit()
    logger.info(
        "product_generation item_code=%s model_source=%s duration_seconds=%.3f retry_count=%s "
        "violation_count=%s status=%s",
        product.product_id,
        model_source,
        process_duration,
        model_retry_count,
        violation_count,
        "FALLBACK_DRAFT" if is_fallback else "PENDING_REVIEW",
    )
    return {
        "thread_id": thread_id,
        "copy": generated,
        "violations": state.get("violations", []),
        "review_status": state.get("review_status"),
        "status": "FALLBACK_DRAFT" if is_fallback else "PENDING_REVIEW",
        "model_source": model_source,
        "retry_count": model_retry_count,
        "process_duration": round(process_duration, 3),
        "error_log": state.get("generation_error"),
    }


@app.post("/api/products/{thread_id}/review")
async def review_product(thread_id: str, review: ProductReview) -> dict[str, Any]:
    if review.action not in {"approve", "approve_with_edits"}:
        raise HTTPException(status_code=422, detail="action 必须为 approve 或 approve_with_edits")
    if review.action == "approve_with_edits" and review.copy is None:
        raise HTTPException(status_code=422, detail="修改后通过时必须提供 copy")
    if review.action == "approve" and review.copy is not None:
        raise HTTPException(status_code=422, detail="approve 操作不接受 copy")

    config = {"configurable": {"thread_id": thread_id}}
    snapshot = await workflow.aget_state(config)
    if not snapshot.values or not snapshot.next:
        raise HTTPException(status_code=404, detail="未找到待审核的商品流程")
    decision = {"action": review.action, "generated_copy": review.copy.model_dump() if review.copy else None}
    try:
        result = await workflow.ainvoke(Command(resume=decision), config=config)
    except (RuntimeError, ValueError) as exc:
        raise HTTPException(status_code=409, detail=f"无法恢复审核流程：{exc}") from exc

    final_copy = result.get("generated_copy")
    if hasattr(final_copy, "model_dump"):
        final_copy = final_copy.model_dump()
    final = {
        "item_code": result.get("item_code"),
        "copy": final_copy,
        "violations": result.get("violations", []),
        "review_status": "APPROVED",
    }
    async with _session_factory()() as session:
        await _save_product(
            session,
            result.get("item_code"),
            status="已通过",
            retry_count=result.get("retry_count", 0),
        )
        if review.action == "approve_with_edits":
            _add_version(session, result["item_code"], "HUMAN_MODIFIED", final_copy)
        await session.commit()
    return final


@app.post("/api/products/generate")
async def generate_product_copy(product: ProductInput) -> dict[str, Any]:
    started_at = time.perf_counter()
    try:
        result = await generate_copy_with_metadata(product)
        process_duration = time.perf_counter() - started_at
        violations = check_violations(result.copy)
        async with _session_factory()() as session:
            await _save_product(
                session,
                product.product_id,
                status=result.status if result.status == "FALLBACK_DRAFT" else "待审核",
                retry_count=result.retry_count,
                error_log=result.error_log,
                process_duration=process_duration,
            )
            _add_version(session, product.product_id, "RAW", product.model_dump(by_alias=True))
            _add_version(
                session,
                product.product_id,
                "FALLBACK_DRAFT" if result.status == "FALLBACK_DRAFT" else "AI_GENERATED",
                result.copy.model_dump(),
            )
            await session.commit()
        logger.info(
            "product_generation item_code=%s model_source=%s duration_seconds=%.3f retry_count=%s "
            "violation_count=%s status=%s",
            product.product_id,
            result.model_source,
            process_duration,
            result.retry_count,
            len(violations),
            result.status,
        )
        return {
            "copy": result.copy,
            "violations": violations,
            "status": result.status,
            "model_source": result.model_source,
            "retry_count": result.retry_count,
            "process_duration": round(process_duration, 3),
            "error_log": result.error_log,
        }
    except RuntimeError as exc:
        await _record_failure(product.product_id, str(exc), time.perf_counter() - started_at)
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except ValidationError as exc:
        raise HTTPException(
            status_code=502,
            detail={"message": "中转站返回的文案不符合 GeneratedCopy 结构", "errors": exc.errors()},
        ) from exc
    except OpenAIError as exc:
        await _record_failure(product.product_id, str(exc), time.perf_counter() - started_at)
        raise HTTPException(status_code=502, detail=f"大模型请求失败：{exc}") from exc

FIELD_NAMES = [
    "产品编号",
    "产品名称",
    "产品特性",
    "产品特性和材质",
    "产品价格",
    "产品的重量",
    "产品的尺寸",
]
FIELD_TO_NAME = {
    "product_id": "产品编号",
    "product_name": "产品名称",
    "product_features": "产品特性",
    "material_and_features": "产品特性和材质",
    "price": "产品价格",
    "weight": "产品的重量",
    "dimensions": "产品的尺寸",
}
KEYWORDS_TO_FIELD = (
    ("材质", "产品特性和材质"),
    ("编号", "产品编号"),
    ("名称", "产品名称"),
    ("价格", "产品价格"),
    ("特性", "产品特性"),
    ("重量", "产品的重量"),
    ("尺寸", "产品的尺寸"),
)


def cell_value(cell: Any) -> str:
    value = cell.value
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header_field(header: str) -> str | None:
    for keyword, field_name in KEYWORDS_TO_FIELD:
        if keyword in header:
            return field_name
    return None


def extract_price(value: str) -> str:
    compact = value.replace(",", "").replace("，", "")
    compact = re.sub(r"\s+", "", compact)
    match = re.search(r"\d+(?:\.\d+)?", compact)
    return match.group(0) if match else ""


@app.post("/api/upload")
async def upload(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".xls", ".xlsx"}:
        raise HTTPException(status_code=400, detail="请上传 .xls 或 .xlsx 格式的商品表格")

    try:
        contents = await file.read()
        if suffix == ".xlsx":
            workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
            worksheet = workbook.active
            rows = [
                ["" if value is None else str(int(value) if isinstance(value, float) and value.is_integer() else value).strip()
                 for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            workbook.close()
        else:
            workbook = xlrd.open_workbook(file_contents=contents)
            sheet = workbook.sheet_by_index(0)
            rows = [
                [cell_value(sheet.cell(row, col)) for col in range(sheet.ncols)]
                for row in range(sheet.nrows)
            ]
    except (xlrd.XLRDError, OSError, ValueError, BadZipFile, KeyError) as exc:
        raise HTTPException(status_code=400, detail=f"无法解析商品表格：{exc}") from exc

    if not rows:
        raise HTTPException(status_code=400, detail="表格为空")

    column_count = max(map(len, rows), default=0)
    headers = [rows[0][col] if col < len(rows[0]) else "" for col in range(column_count)]
    column_by_name: dict[str, int] = {}
    for col, header in enumerate(headers):
        field_name = header_field(header)
        if field_name and field_name not in column_by_name:
            column_by_name[field_name] = col

    # Use the conventional seven-column order for fields whose headers are unrecognized.
    used_columns = set(column_by_name.values())
    for position, field_name in enumerate(FIELD_NAMES):
        if field_name not in column_by_name and position < column_count and position not in used_columns:
            column_by_name[field_name] = position
            used_columns.add(position)
    for field_name in FIELD_NAMES:
        if field_name not in column_by_name and column_count:
            column_by_name[field_name] = min(len(column_by_name), column_count - 1)

    imported: list[ProductInput] = []
    errors: list[dict[str, Any]] = []
    for row_index in range(1, len(rows)):
        row_number = row_index + 1
        product_id_col = column_by_name["产品编号"]
        source_row = rows[row_index]
        product_id = source_row[product_id_col] if product_id_col < len(source_row) else ""
        if not product_id:
            continue

        values = {
            field_name: source_row[column] if column < len(source_row) else ""
            for field_name, column in column_by_name.items()
        }
        price = extract_price(values["产品价格"])
        if not price:
            price_col = column_by_name["产品价格"]
            adjacent_columns = [col for col in (price_col - 1, price_col + 1) if 0 <= col < column_count]
            adjacent_values = [source_row[col] if col < len(source_row) else "" for col in adjacent_columns]
            numeric_values = [extract_price(value) for value in adjacent_values if re.search(r"\d", value)]
            plain_numeric_values = [
                extract_price(value)
                for value in adjacent_values
                if re.fullmatch(r"\s*[￥¥]?\s*[\d\s,.]+\s*(?:元)?\s*", value)
            ]
            price = (plain_numeric_values or numeric_values or [""])[0]
        values["产品价格"] = price
        try:
            imported.append(ProductInput.model_validate(values))
        except ValidationError as exc:
            reasons = [
                f"{error['loc'][0]}：{error['msg']}" for error in exc.errors()
            ]
            errors.append({"row": row_number, "reason": reasons})

    if imported:
        try:
            async with _session_factory()() as session:
                for product in imported:
                    await _save_product(session, product.product_id, status="RAW", retry_count=0)
                    _add_version(session, product.product_id, "RAW", product.model_dump(by_alias=True))
                await session.commit()
        except Exception as exc:
            raise HTTPException(status_code=503, detail=f"商品写入数据库失败：{exc}") from exc
    products.extend(imported)
    task_id = str(uuid.uuid4())
    task_progress[task_id] = {
        "status": "processing" if imported else "completed",
        "total_products": len(imported),
        "processed_products": 0,
        "processed_count": 0,
        "success_count": 0,
        "products": [
            {
                "item_code": product.product_id,
                "raw": product.model_dump(by_alias=True),
                "copy": None,
                "violations": [],
                "thread_id": None,
                "status": "RAW",
            }
            for product in imported
        ],
        "updates": [],
        "errors": [],
        "validation_errors": errors,
    }
    task_controls[task_id] = asyncio.Event()
    task_controls[task_id].set()
    if imported:
        background_tasks.add_task(_process_upload_task, task_id, imported)
    return {"task_id": task_id, "status": task_progress[task_id]["status"]}


@app.get("/api/export")
async def export() -> StreamingResponse:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("商品数据")
    for col, name in enumerate(FIELD_NAMES):
        sheet.write(0, col, name)
    for row, product in enumerate(products, start=1):
        values = product.model_dump(by_alias=True)
        for col, name in enumerate(FIELD_NAMES):
            sheet.write(row, col, values[name])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.ms-excel",
        headers={"Content-Disposition": 'attachment; filename="products.xls"'},
    )
